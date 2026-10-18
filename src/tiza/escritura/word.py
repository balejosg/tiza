"""Escribe un Word sencillo (.docx) a partir de Markdown, sin salirse de lo admitido.

Admite títulos 1-4, párrafos, negrita, cursiva, código en línea (como texto), enlaces
(«texto (URL)»), listas con viñetas y numeradas hasta 3 niveles, tablas GFM y saltos de página
(una línea ``\\pagebreak``). Rechaza el HTML, las imágenes, los bloques de código y las citas
con ``MARKDOWN_NO_ADMITIDO`` y la línea: así no hay rutas locales que resolver ni nada que
ejecutar. El fichero se construye entero en memoria y nunca sustituye a uno existente.
"""

from __future__ import annotations

import io
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from markdown_it import MarkdownIt
from markdown_it.token import Token

from tiza.escritura import ErrorEscritura

__all__ = ["escribir_docx"]

MAX_NIVEL_TITULO = 4
MAX_NIVEL_LISTA = 3
_SALTO_DE_PAGINA = "\\pagebreak"
_CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_NO_ADMITIDO = {
    "html_block": "el HTML incrustado",
    "html_inline": "el HTML incrustado",
    "image": "las imágenes",
    "code_block": "los bloques de código",
    "fence": "los bloques de código",
    "blockquote_open": "las citas",
}


def _rechazo(linea: int | None, que: str) -> ErrorEscritura:
    donde = f"línea {linea}" if linea else "línea desconocida"
    return ErrorEscritura("MARKDOWN_NO_ADMITIDO", f"{donde}: no se admite {que}.")


def _linea(token: Token) -> int | None:
    return token.map[0] + 1 if token.map else None


def _validar(tokens: list[Token]) -> None:
    """Rechaza, antes de escribir nada, todo lo que el Word sencillo no admite."""
    nivel_lista = 0
    for token in tokens:
        if token.type in _NO_ADMITIDO:
            raise _rechazo(_linea(token), _NO_ADMITIDO[token.type])
        if token.type == "heading_open" and int(token.tag[1]) > MAX_NIVEL_TITULO:
            raise _rechazo(_linea(token), f"los títulos de nivel {token.tag[1]}")
        if token.type in ("bullet_list_open", "ordered_list_open"):
            nivel_lista += 1
            if nivel_lista > MAX_NIVEL_LISTA:
                raise _rechazo(_linea(token), "listas anidadas a más de 3 niveles")
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            nivel_lista -= 1
        elif token.type == "inline":
            for hijo in token.children or []:
                if hijo.type in _NO_ADMITIDO:
                    raise _rechazo(_linea(token), _NO_ADMITIDO[hijo.type])
            if _CONTROL.search(token.content):
                raise _rechazo(_linea(token), "caracteres de control")


IDIOMA = "es-ES"
_LANG = "w:docDefaults/w:rPrDefault/w:rPr/w:lang"


def _plantilla_en_espanol() -> Any:
    """El documento base: idioma español y sin la firma de la librería en las propiedades."""
    documento = Document()
    for lang in documento.styles.element.xpath(_LANG):
        for atributo in ("w:val", "w:eastAsia", "w:bidi"):
            if lang.get(qn(atributo)) is not None:
                lang.set(qn(atributo), IDIOMA)
    ahora = datetime.now(timezone.utc).replace(microsecond=0)
    propiedades = documento.core_properties
    propiedades.author = "tiza"
    propiedades.last_modified_by = "tiza"
    propiedades.created = ahora
    propiedades.modified = ahora
    propiedades.comments = ""
    propiedades.title = ""
    propiedades.subject = ""
    propiedades.keywords = ""
    return documento


class _Constructor:
    def __init__(self) -> None:
        self.documento = _plantilla_en_espanol()
        self.bloques = 0

    def _inline(self, parrafo: Any, token: Token, negrita: bool = False) -> None:
        negritas = int(negrita)
        cursivas = 0
        enlaces: list[list[str]] = []  # [href, texto escrito dentro del enlace]

        def escribir(texto: str, salto: bool = False) -> None:
            run = parrafo.add_run(texto)
            if salto:
                run.add_break()
            if negritas:
                run.bold = True
            if cursivas:
                run.italic = True
            if enlaces:
                enlaces[-1][1] += texto

        for hijo in token.children or []:
            if hijo.type in ("text", "code_inline"):
                escribir(hijo.content)
            elif hijo.type == "softbreak":
                escribir(" ")
            elif hijo.type == "hardbreak":
                escribir("", salto=True)
            elif hijo.type == "strong_open":
                negritas += 1
            elif hijo.type == "strong_close":
                negritas -= 1
            elif hijo.type == "em_open":
                cursivas += 1
            elif hijo.type == "em_close":
                cursivas -= 1
            elif hijo.type == "link_open":
                enlaces.append([str(hijo.attrGet("href") or ""), ""])
            elif hijo.type == "link_close":
                href, texto = enlaces.pop()
                if href and href != texto:
                    escribir(f" ({href})")

    def _numeracion_nueva(self, estilo: str, inicio: int) -> int:
        """Un ``w:num`` propio, con la misma definición que el estilo, que empieza en ``inicio``.

        Todos los párrafos de un estilo comparten ``numId``: sin esto, una segunda lista
        numerada seguiría contando donde acabó la primera.
        """
        num_id = self.documento.styles[estilo].element.pPr.numPr.numId.val
        numeracion = self.documento.part.numbering_part.element
        abstracto = numeracion.num_having_numId(num_id).abstractNumId.val
        nuevo = numeracion.add_num(abstracto)
        nuevo.add_lvlOverride(0).add_startOverride(inicio)
        return int(nuevo.numId)

    def _tabla(self, tokens: list[Token], i: int) -> int:
        filas: list[list[Token]] = []
        while tokens[i].type != "table_close":
            if tokens[i].type == "tr_open":
                filas.append([])
            elif tokens[i].type in ("th_open", "td_open"):
                filas[-1].append(tokens[i + 1])
            i += 1
        columnas = max(len(fila) for fila in filas)
        tabla = self.documento.add_table(rows=len(filas), cols=columnas)
        tabla.style = "Table Grid"
        for n, fila in enumerate(filas):
            for j, inline in enumerate(fila):
                self._inline(tabla.cell(n, j).paragraphs[0], inline, negrita=n == 0)
        self.bloques += 1
        return i

    def _parrafo(self, listas: list[tuple[str, int | None]]) -> Any:
        if not listas:
            return self.documento.add_paragraph()
        estilo, num_id = listas[-1]
        parrafo = self.documento.add_paragraph(style=estilo)
        if num_id is not None:
            parrafo._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_id
        return parrafo

    def construir(self, tokens: list[Token]) -> None:
        listas: list[tuple[str, int | None]] = []  # (estilo, numId propio o None) por nivel
        i = 0
        while i < len(tokens):
            token = tokens[i]
            if token.type == "heading_open":
                parrafo = self.documento.add_paragraph(style=f"Heading {token.tag[1]}")
                self._inline(parrafo, tokens[i + 1])
                self.bloques += 1
                i += 2
            elif token.type == "paragraph_open":
                inline = tokens[i + 1]
                if not listas and inline.content.strip() == _SALTO_DE_PAGINA:
                    self.documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                else:
                    self._inline(self._parrafo(listas), inline)
                self.bloques += 1
                i += 2
            elif token.type in ("bullet_list_open", "ordered_list_open"):
                nivel = len(listas) + 1
                base = "List Bullet" if token.type == "bullet_list_open" else "List Number"
                estilo = base if nivel == 1 else f"{base} {nivel}"
                num_id = None
                if base == "List Number":
                    num_id = self._numeracion_nueva(estilo, int(token.attrGet("start") or 1))
                listas.append((estilo, num_id))
            elif token.type in ("bullet_list_close", "ordered_list_close"):
                listas.pop()
            elif token.type == "table_open":
                i = self._tabla(tokens, i)
            i += 1


def _escribir_sin_sobrescribir(datos: bytes, destino: Path) -> None:
    """Crea ``destino`` en exclusiva y lo escribe; si algo falla, no deja nada."""
    try:
        fichero = open(destino, "xb")  # noqa: SIM115 - se cierra en el ``with`` siguiente
    except FileExistsError:
        raise ErrorEscritura(
            "DOCUMENTO_EXISTE", f"Ya existe {destino.name}: tiza no sobrescribe un Word."
        ) from None
    except OSError:
        raise ErrorEscritura(
            "DOCUMENTO_NO_ESCRIBIBLE", f"No se pudo crear {destino.name}."
        ) from None
    try:
        with fichero:
            fichero.write(datos)
            fichero.flush()
            os.fsync(fichero.fileno())
    except BaseException as error:
        try:
            destino.unlink()
        except OSError:
            pass
        if isinstance(error, OSError):
            raise ErrorEscritura(
                "DOCUMENTO_NO_ESCRIBIBLE", f"No se pudo escribir {destino.name}."
            ) from None
        raise


def escribir_docx(markdown: str, destino: Path) -> int:
    """Escribe el Word de ``markdown`` en ``destino`` y devuelve los bloques escritos.

    Nunca sustituye un fichero existente (``DOCUMENTO_EXISTE``). Si algo falla a mitad no
    queda nada escrito.
    """
    tokens = MarkdownIt("commonmark").enable("table").parse(markdown)
    _validar(tokens)
    constructor = _Constructor()
    constructor.construir(tokens)
    memoria = io.BytesIO()
    constructor.documento.save(memoria)
    _escribir_sin_sobrescribir(memoria.getvalue(), Path(destino))
    return constructor.bloques
