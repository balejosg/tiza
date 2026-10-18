"""Escritura de un Word sencillo desde Markdown (tiza documento)."""

from __future__ import annotations

import hashlib
import json

import pytest
from docx import Document
from docx.oxml.ns import qn

from tiza import agente, ayuda, cli, informe
from tiza.escritura import ErrorEscritura
from tiza.escritura.word import escribir_docx
from tiza.lectura.word import leer_docx


def test_documento_titulos_y_tabla(tmp_path):
    md = "# Unidad 4\n\nUn párrafo.\n\n| A | B | C |\n|---|---|---|\n| 1 | 2 | 3 |\n"
    destino = tmp_path / "u.docx"
    assert escribir_docx(md, destino) == 3
    bloques = leer_docx(destino).bloques
    assert (bloques[0].tipo, bloques[0].nivel, bloques[0].texto) == ("titulo", 1, "Unidad 4")
    assert bloques[1].texto == "Un párrafo."
    tabla = bloques[2]
    assert (tabla.tipo, tabla.ancho, tabla.alto) == ("tabla", 3, 2)
    textos = [[c.hijos[0].texto for c in fila.hijos] for fila in tabla.hijos]
    assert textos == [["A", "B", "C"], ["1", "2", "3"]]
    primera = Document(str(destino)).tables[0]
    assert primera.style.name == "Table Grid"
    assert all(r.bold for c in primera.rows[0].cells for r in c.paragraphs[0].runs)
    assert not any(r.bold for c in primera.rows[1].cells for r in c.paragraphs[0].runs)


def test_documento_listas_numeradas(tmp_path):
    destino = tmp_path / "u.docx"
    escribir_docx("1. Uno\n2. Dos\n", destino)
    assert [(b.etiqueta, b.texto) for b in leer_docx(destino).bloques] == [
        ("1.", "Uno"),
        ("2.", "Dos"),
    ]


def test_dos_listas_numeradas_reinician(tmp_path):
    destino = tmp_path / "u.docx"
    escribir_docx("1. Uno\n2. Dos\n\nEntre medias.\n\n1. Tres\n2. Cuatro\n", destino)
    etiquetas = [b.etiqueta for b in leer_docx(destino).bloques]
    assert etiquetas == ["1.", "2.", "", "1.", "2."]


def test_listas_anidadas_y_viñetas(tmp_path):
    destino = tmp_path / "u.docx"
    escribir_docx("- a\n  - b\n    - c\n- d\n\n1. x\n   1. y\n", destino)
    bloques = leer_docx(destino).bloques
    assert [(b.estilo, b.texto) for b in bloques] == [
        ("List Bullet", "a"),
        ("List Bullet 2", "b"),
        ("List Bullet 3", "c"),
        ("List Bullet", "d"),
        ("List Number", "x"),
        ("List Number 2", "y"),
    ]
    assert [b.etiqueta for b in bloques][-2:] == ["1.", "1."]


def test_lista_anidada_demasiado_profunda(tmp_path):
    md = "- a\n  - b\n    - c\n      - d\n"
    with pytest.raises(ErrorEscritura) as exc:
        escribir_docx(md, tmp_path / "u.docx")
    assert exc.value.codigo == "MARKDOWN_NO_ADMITIDO" and "línea 4" in exc.value.detalle
    assert not (tmp_path / "u.docx").exists()


def test_formato_en_linea_y_enlaces(tmp_path):
    destino = tmp_path / "u.docx"
    escribir_docx("Hola **negrita**, *cursiva* y `código` en [web](https://x.es/a).\n", destino)
    p = Document(str(destino)).paragraphs[0]
    assert p.text == "Hola negrita, cursiva y código en web (https://x.es/a)."
    assert [r.text for r in p.runs if r.bold] == ["negrita"]
    assert [r.text for r in p.runs if r.italic] == ["cursiva"]


def test_salto_de_pagina_y_regla(tmp_path):
    destino = tmp_path / "u.docx"
    assert escribir_docx("Uno\n\n\\pagebreak\n\n---\n\nDos\n", destino) == 3
    xml = Document(str(destino)).element.xml
    assert 'w:br w:type="page"' in xml
    assert "pagebreak" not in xml


def test_titulo_nivel_cinco_rechazado(tmp_path):
    with pytest.raises(ErrorEscritura) as exc:
        escribir_docx("Hola\n\n##### Cinco\n", tmp_path / "u.docx")
    assert exc.value.codigo == "MARKDOWN_NO_ADMITIDO" and "línea 3" in exc.value.detalle


@pytest.mark.parametrize(
    "md",
    [
        "Hola\n\n<script>alert(1)</script>\n",
        "Hola <b>x</b>\n",
        "![alt](a.png)\n",
        "    sangrado\n",
        "```\ncodigo\n```\n",
        "> cita\n",
    ],
)
def test_documento_rechaza_lo_no_admitido(tmp_path, md):
    with pytest.raises(ErrorEscritura) as exc:
        escribir_docx(md, tmp_path / "u.docx")
    assert exc.value.codigo == "MARKDOWN_NO_ADMITIDO"
    assert exc.value.detalle.startswith("línea ")


def test_documento_rechaza_html(tmp_path):
    with pytest.raises(ErrorEscritura) as exc:
        escribir_docx("Uno\n\n<script>x</script>\n", tmp_path / "u.docx")
    assert exc.value.codigo == "MARKDOWN_NO_ADMITIDO"
    assert "línea 3" in exc.value.detalle


def test_documento_no_sobrescribe(tmp_path):
    destino = tmp_path / "prog.docx"
    destino.write_bytes(b"original")
    antes = hashlib.sha256(destino.read_bytes()).hexdigest()
    with pytest.raises(ErrorEscritura) as exc:
        escribir_docx("# Hola\n", destino)
    assert exc.value.codigo == "DOCUMENTO_EXISTE"
    assert hashlib.sha256(destino.read_bytes()).hexdigest() == antes


def test_documento_atomico(tmp_path, monkeypatch):
    destino = tmp_path / "u.docx"

    def falla(_descriptor):
        raise OSError("disco lleno")

    monkeypatch.setattr("tiza.escritura.word.os.fsync", falla)
    with pytest.raises(ErrorEscritura) as exc:
        escribir_docx("# Hola\n", destino)
    assert exc.value.codigo == "DOCUMENTO_NO_ESCRIBIBLE"
    assert list(tmp_path.iterdir()) == []


def test_codigos_nuevos_tienen_ayuda():
    for codigo in ("MARKDOWN_NO_ADMITIDO", "DOCUMENTO_EXISTE", "DOCUMENTO_NO_ESCRIBIBLE"):
        assert codigo in ayuda.MENSAJES
    assert "--salida" in ayuda.MENSAJES["DOCUMENTO_EXISTE"]


def test_resumen_esquema_cerrado():
    resumen = {"fuente": "a.md", "documento": "a.docx", "bloques": 2}
    assert informe.Documento(**resumen).a_dict() == resumen
    with pytest.raises(informe.ErrorInforme) as exc:
        informe.validar_documento({**resumen, "extra": "x"})
    assert exc.value.codigo == "CAMPO_NO_PERMITIDO"
    with pytest.raises(informe.ErrorInforme):
        informe.validar_documento({**resumen, "bloques": "2"})
    with pytest.raises(informe.ErrorInforme):
        informe.validar_documento({**resumen, "fuente": "a<b>.md"})


# --- Agente y CLI ---------------------------------------------------------------------------


@pytest.fixture
def carpeta(tmp_path, monkeypatch):
    asignatura = tmp_path / "asignatura"
    asignatura.mkdir()
    monkeypatch.chdir(asignatura)
    return asignatura


def ejecutar(capsys, argv):
    codigo = cli.main(argv)
    return codigo, json.loads(capsys.readouterr().out)


def test_cli_documento_por_defecto(carpeta, capsys):
    (carpeta / "u.md").write_text("# Unidad\n\nTexto\n", encoding="utf-8")
    codigo, resumen = ejecutar(capsys, ["documento", "u.md"])
    assert codigo == 0
    assert resumen == {"fuente": "u.md", "documento": "u.docx", "bloques": 2}
    assert (carpeta / "u.docx").is_file()
    assert not (carpeta / ".tiza" / "informe.json").exists()
    assert "documento" in agente.__all__


def test_cli_documento_con_salida(carpeta, capsys):
    (carpeta / "u.md").write_text("Texto\n", encoding="utf-8")
    codigo, resumen = ejecutar(capsys, ["documento", "u.md", "--salida", "sub/x.docx"])
    assert codigo == 2 and resumen["error"] == "DOCUMENTO_NO_ESCRIBIBLE"  # sub no existe
    (carpeta / "sub").mkdir()
    codigo, resumen = ejecutar(capsys, ["documento", "u.md", "--salida", "sub/x.docx"])
    assert codigo == 0 and resumen["documento"] == "sub/x.docx"


@pytest.mark.parametrize(
    ("argv", "error"),
    [
        (["documento", "../fuera.md"], "FICHERO_FUERA"),
        (["documento", "ausente.md"], "DOCUMENTO_AUSENTE"),
        (["documento", "u.txt"], "FORMATO_NO_ADMITIDO"),
        (["documento", "u.md", "--salida", "x.txt"], "FORMATO_NO_ADMITIDO"),
        (["documento", "u.md", "--salida", ".oculto.docx"], "FICHERO_TEMPORAL"),
        (["documento", "u.md", "--salida", "../x.docx"], "FICHERO_FUERA"),
        (["documento", ".tiza/u.md"], "FICHERO_TEMPORAL"),
    ],
)
def test_cli_documento_errores(carpeta, capsys, argv, error):
    (carpeta / "u.md").write_text("Texto\n", encoding="utf-8")
    codigo, salida = ejecutar(capsys, argv)
    assert codigo == 2 and salida["error"] == error
    assert not list(carpeta.glob("*.docx"))


def test_cli_documento_no_utf8(carpeta, capsys):
    (carpeta / "u.md").write_bytes(b"\xff\xfe\x00 mal")
    codigo, salida = ejecutar(capsys, ["documento", "u.md"])
    assert codigo == 2 and salida["error"] == "FORMATO_NO_ADMITIDO"


def test_cli_documento_demasiado_grande(carpeta, capsys):
    (carpeta / "u.md").write_bytes(b"a" * (2 * 1024 * 1024 + 1))
    codigo, salida = ejecutar(capsys, ["documento", "u.md"])
    assert codigo == 2 and salida["error"] == "FICHERO_DEMASIADO_GRANDE"


def test_cli_documento_existe(carpeta, capsys):
    (carpeta / "u.md").write_text("Texto\n", encoding="utf-8")
    (carpeta / "u.docx").write_bytes(b"mio")
    codigo, salida = ejecutar(capsys, ["documento", "u.md"])
    assert codigo == 2 and salida["error"] == "DOCUMENTO_EXISTE"
    assert (carpeta / "u.docx").read_bytes() == b"mio"


def test_cli_documento_markdown_no_admitido(carpeta, capsys):
    (carpeta / "u.md").write_text("Hola\n\n<div>x</div>\n", encoding="utf-8")
    codigo, salida = ejecutar(capsys, ["documento", "u.md"])
    assert codigo == 2 and salida["error"] == "MARKDOWN_NO_ADMITIDO"
    assert "línea 3" in salida["detalle"]
    assert not (carpeta / "u.docx").exists()


def test_word_en_espanol_y_sin_firma_de_python_docx(tmp_path):
    import zipfile

    destino = tmp_path / "u.docx"
    escribir_docx("# Unidad\n", destino)
    documento = Document(str(destino))
    lang = documento.styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr/w:lang")[0]
    assert lang.get(qn("w:val")) == "es-ES"
    assert lang.get(qn("w:eastAsia")) == "es-ES"
    propiedades = documento.core_properties
    assert propiedades.author == "tiza"
    assert propiedades.last_modified_by == "tiza"
    assert propiedades.created.year >= 2026
    assert propiedades.modified.year >= 2026
    assert (propiedades.comments, propiedades.title) == ("", "")
    assert (propiedades.subject, propiedades.keywords) == ("", "")
    with zipfile.ZipFile(destino) as z:
        assert not any(b"python-docx" in z.read(n) for n in z.namelist())
