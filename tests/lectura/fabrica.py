"""Construye los .docx sintéticos de los tests con python-docx y XML insertado."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.oxml import parse_xml

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def elemento(texto: str):
    """Parsea un elemento w:* suelto (declara el espacio de nombres en la raíz)."""
    raiz, resto = texto.split(">", 1)
    if raiz.endswith("/"):
        return parse_xml(f"{raiz[:-1]} {W}/>")
    return parse_xml(f"{raiz} {W}>{resto}")


def nuevo():
    return Document()


def parrafo_xml(documento, interior: str):
    """Añade un párrafo cuyo contenido es XML en bruto (w:r, w:ins, w:del…)."""
    p = documento.add_paragraph()
    for hijo in elemento(f"<w:x>{interior}</w:x>"):
        p._p.append(hijo)
    return p


def indice(documento) -> None:
    """Un índice de contenidos (w:sdt con docPartGallery) al principio del cuerpo."""
    sdt = elemento(
        '<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
        "</w:docPartObj></w:sdtPr><w:sdtContent><w:p><w:r><w:t>Entrada del índice</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )
    documento.element.body.insert(0, sdt)


def nivel(ilvl: int, fmt: str, texto: str, inicio: int = 1, estilo: str = "") -> str:
    ps = f'<w:pStyle w:val="{estilo}"/>' if estilo else ""
    return (
        f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{inicio}"/><w:numFmt w:val="{fmt}"/>'
        f'{ps}<w:lvlText w:val="{texto}"/></w:lvl>'
    )


def numeracion(
    documento, abstractos: dict[int, list[str]], nums: dict[int, tuple[int, int | None]]
):
    """Sustituye numbering.xml. ``nums``: numId -> (abstractNumId, startOverride del nivel 0)."""
    partes = [f"<w:numbering {W}>"]
    for ident, niveles in abstractos.items():
        partes.append(
            f'<w:abstractNum w:abstractNumId="{ident}">{"".join(niveles)}</w:abstractNum>'
        )
    for num, (abstracto, override) in nums.items():
        extra = (
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{override}"/></w:lvlOverride>'
            if override
            else ""
        )
        partes.append(
            f'<w:num w:numId="{num}"><w:abstractNumId w:val="{abstracto}"/>{extra}</w:num>'
        )
    partes.append("</w:numbering>")
    documento.part.numbering_part._element = parse_xml("".join(partes))


def parrafo_numerado(documento, texto: str, num: int, ilvl: int = 0):
    p = documento.add_paragraph(texto)
    p._p.get_or_add_pPr().append(
        elemento(f'<w:numPr><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num}"/></w:numPr>')
    )
    return p


def estilo_titulo_en_espanol(documento) -> None:
    """Un estilo «Título 2» con el ``styleId`` que le pone Word en español (``Ttulo2``)."""
    documento.styles.element.append(
        elemento(
            '<w:style w:type="paragraph" w:styleId="Ttulo2"><w:name w:val="Título 2"/>'
            '<w:basedOn w:val="Normal"/><w:pPr><w:outlineLvl w:val="1"/></w:pPr></w:style>'
        )
    )


def tabla_con_total(documento):
    """Tabla RA | texto | % con cabecera, dos filas y un total fusionado (``gridSpan`` 2)."""
    tabla = documento.add_table(rows=4, cols=3)
    filas = [("RA", "Resultado de aprendizaje", "%"), ("RA1", "Reconoce A", "40"),
             ("RA2", "Reconoce B", "60")]  # fmt: skip
    for n, fila in enumerate(filas):
        for j, texto in enumerate(fila):
            tabla.cell(n, j).text = texto
    tabla.cell(3, 0).merge(tabla.cell(3, 1)).text = "TOTAL"
    tabla.cell(3, 2).text = "100"
    return tabla


def tabla_vmerge(documento):
    """Tabla de 3×2 cuya primera columna está fusionada en vertical en las dos primeras filas."""
    tabla = documento.add_table(rows=3, cols=2)
    tabla.cell(0, 0).merge(tabla.cell(1, 0)).text = "Bloque"
    tabla.cell(0, 1).text = "Uno"
    tabla.cell(1, 1).text = "Dos"
    tabla.cell(2, 0).text = "Suelta"
    tabla.cell(2, 1).text = "Tres"
    return tabla


def tabla_unidad(documento):
    """La tabla de una unidad de trabajo: título fusionado en la primera fila y dos filas más."""
    tabla = documento.add_table(rows=3, cols=2)
    tabla.cell(0, 0).merge(tabla.cell(0, 1)).text = "UNIDAD DE TRABAJO 1"
    for n, (a, b) in enumerate([("Título", "Introducción"), ("Horas", "12")], start=1):
        tabla.cell(n, 0).text = a
        tabla.cell(n, 1).text = b
    return tabla


def encabezado_y_pie(documento, cabecera: str, pie: str) -> None:
    seccion = documento.sections[0]
    seccion.header.paragraphs[0].text = cabecera
    seccion.footer.paragraphs[0].text = pie


def pie_con_pagina(seccion, antes: str = "Página ") -> None:
    """El pie de la sección: «<antes>» seguido del campo PAGE (con su resultado en caché)."""
    pie = seccion.footer
    pie.is_linked_to_previous = False
    p = pie.paragraphs[0]
    p.text = antes
    for hijo in elemento(
        '<w:x><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:x>'
    ):
        p._p.append(hijo)


def dos_secciones(documento, cabecera: str = "Curso 2026/27") -> None:
    """Dos secciones que comparten encabezado, cada una con su pie (el 1.º con el campo PAGE).

    El cuerpo queda con dos párrafos antes del salto de sección y uno después.
    """
    documento.add_paragraph("Antes 1")
    documento.add_paragraph("Antes 2")
    documento.add_section()
    documento.add_paragraph("Después")
    primera, segunda = documento.sections
    primera.header.paragraphs[0].text = cabecera
    pie_con_pagina(primera)
    segunda.header.is_linked_to_previous = True  # comparten el de la primera
    segunda.footer.is_linked_to_previous = False
    segunda.footer.paragraphs[0].text = "Departamento de Informática"


def lista_de_letras(documento, textos: list[str]) -> None:
    """Una lista numerada a), b), c)… (el ``numId`` es el 1)."""
    numeracion(documento, {1: [nivel(0, "lowerLetter", "%1)")]}, {1: (1, None)})
    for texto in textos:
        parrafo_numerado(documento, texto, 1)


def control_en_linea(documento, marcador: str = "Haga clic aquí"):
    """Un párrafo con un control de contenido en línea que enseña su texto de ejemplo."""
    return parrafo_xml(
        documento,
        '<w:sdt><w:sdtPr><w:alias w:val="Asignatura"/><w:showingPlcHdr/></w:sdtPr><w:sdtContent>'
        '<w:r><w:rPr><w:rStyle w:val="PlaceholderText"/><w:b/></w:rPr>'
        f"<w:t>{marcador}</w:t></w:r></w:sdtContent></w:sdt>",
    )


def a_bytes(documento) -> bytes:
    memoria = io.BytesIO()
    documento.save(memoria)
    return memoria.getvalue()


def reescribir(datos: bytes, miembro: str, contenido: bytes) -> bytes:
    """Copia un zip cambiando (o añadiendo) un miembro."""
    salida = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(datos)) as origen,
        zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino,
    ):
        for info in origen.infolist():
            if info.filename != miembro:
                destino.writestr(info.filename, origen.read(info.filename))
        destino.writestr(miembro, contenido)
    return salida.getvalue()


def _escapar_pdf(linea: str) -> str:
    return linea.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_bytes(paginas: list[list[str]]) -> bytes:
    """Un PDF mínimo válido (Helvetica, xref correcta). Cada página es una lista de líneas;
    una lista vacía da una página con el flujo de contenido vacío."""
    n = len(paginas)
    objetos: list[bytes] = [b""] * (3 + 2 * n)
    objetos[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    hijos = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objetos[1] = f"<< /Type /Pages /Kids [{hijos}] /Count {n} >>".encode()
    objetos[2] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i, lineas in enumerate(paginas):
        partes = []
        for k, linea in enumerate(lineas):
            partes.append(f"BT /F1 12 Tf 50 {750 - 20 * k} Td ({_escapar_pdf(linea)}) Tj ET")
        flujo = "\n".join(partes).encode("latin-1")
        objetos[3 + 2 * i] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 3 0 R >> >> " + f"/Contents {5 + 2 * i} 0 R >>".encode()
        )
        objetos[4 + 2 * i] = (
            f"<< /Length {len(flujo)} >>\nstream\n".encode() + flujo + b"\nendstream"
        )
    salida = bytearray(b"%PDF-1.4\n")
    posiciones = []
    for numero, cuerpo in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += f"{numero} 0 obj\n".encode() + cuerpo + b"\nendobj\n"
    inicio = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n".encode() + b"0000000000 65535 f \n"
    for pos in posiciones:
        salida += f"{pos:010d} 00000 n \n".encode()
    salida += (
        f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{inicio}\n%%EOF\n"
    ).encode()
    return bytes(salida)


def cifrar_pdf(datos: bytes, clave: str) -> bytes:
    """El mismo PDF cifrado con RC4-128 y la contraseña de usuario ``clave``."""
    from pypdf import PdfReader, PdfWriter

    escritor = PdfWriter(clone_from=PdfReader(io.BytesIO(datos)))
    escritor.encrypt(user_password=clave, algorithm="RC4-128")
    salida = io.BytesIO()
    escritor.write(salida)
    return salida.getvalue()
